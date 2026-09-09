from pathlib import Path
from zipfile import BadZipFile, ZipFile

import openpyxl
from django.core.exceptions import ValidationError
from django.db import transaction

from ..models import VictimProfile

COLUMN_FIELDS = {2: 'profile_number', 3: 'pcz_ivena', 4: 'category', 5: 'expected_med_action', 6: 'diagnosis', 7: 'visual_diagnosis', 8: 'findings', 9: 'symptoms', 10: 'actor_hints', 11: 'required_specialty', 12: 'gcs', 13: 'spo2', 14: 'rekap', 15: 'resp_rate', 16: 'sys_rr', 17: 'ekg_monitor', 18: 'ro_thorax', 19: 'fast_sono', 20: 'e_fast', 21: 'radiology_finds', 22: 'hb_value', 23: 'blood_units', 24: 'red_treatment_area', 25: 'ventilation_place', 26: 'icu_place', 27: 'emergency_op', 28: 'op_sieve_special', 29: 'op_sieve_basic', 30: 'personal_resources', 31: 'anesthesia_team', 32: 'radiology_resources', 33: 'op_achi_res', 34: 'op_uchi_res', 35: 'op_nchi_res', 36: 'medications', 37: 'pre_treatment_rd', 38: 'spare_col1', 39: 'spare_col2', 46: 'scenario_field', 47: 'comment', 48: 'lastname', 49: 'firstname', 50: 'birthdate'}


def read_profiles(file):
    if Path(file.name).suffix.lower() != '.xlsx' or file.size > 10 * 1024 * 1024:
        raise ValidationError('Bitte eine XLSX-Datei mit höchstens 10 MB hochladen.')
    try:
        file.seek(0)
        with ZipFile(file) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValidationError('Die entpackte Excel-Datei ist zu groß.')
        file.seek(0)
        workbook = openpyxl.load_workbook(file, read_only=True, data_only=False)
        try:
            sheet = workbook.active
            if sheet.max_row > 10001 or sheet.max_column > 100:
                raise ValidationError('Maximal 10000 Datenzeilen und 100 Spalten sind erlaubt.')
            rows = []
            seen = set()
            for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=51), start=2):
                if cells[2].value in (None, ''):
                    continue
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValidationError('Formeln sind in Importdateien nicht erlaubt. Bitte Werte einfügen.')
                values = {field: str(cells[column].value).strip() if cells[column].value is not None else '' for column, field in COLUMN_FIELDS.items()}
                if values['profile_number'] in seen:
                    raise ValidationError('Profilnummern dürfen in der Datei nicht mehrfach vorkommen.')
                seen.add(values['profile_number'])
                profile = VictimProfile(**values)
                try:
                    profile.full_clean(validate_unique=False)
                except ValidationError as exc:
                    details = '; '.join(
                        f'{VictimProfile._meta.get_field(field).verbose_name}: {message}'
                        for field, errors in exc.message_dict.items() for message in errors
                    )
                    raise ValidationError(f'Zeile {row_number}: {details}') from exc
                rows.append(values)
            if not rows:
                raise ValidationError('Die Datei enthält keine Profile in Spalte C.')
            return rows
        finally:
            workbook.close()
    except (BadZipFile, OSError, KeyError, ValueError) as exc:
        raise ValidationError('Die Excel-Datei ist ungültig.') from exc
    finally:
        file.seek(0)


@transaction.atomic
def import_profiles(rows):
    for values in rows:
        values = dict(values)
        number = values.pop('profile_number')
        VictimProfile.objects.update_or_create(profile_number=number, defaults=values)
    return len(rows)
